"""Design system and drawing helpers for the Burj Khalifa deck."""

from __future__ import annotations

import math

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_MARK,
)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- palette
SW, SH = 13.333, 7.5
M = 0.7

BG = RGBColor(0x08, 0x0E, 0x19)
BG2 = RGBColor(0x0B, 0x14, 0x24)
PANEL = RGBColor(0x10, 0x1A, 0x2C)
PANEL2 = RGBColor(0x15, 0x22, 0x38)
GRID = RGBColor(0x22, 0x31, 0x4A)
TEXT = RGBColor(0xE8, 0xEE, 0xF7)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
DIM = RGBColor(0x64, 0x74, 0x8B)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)
RED = RGBColor(0xF8, 0x71, 0x71)
GREEN = RGBColor(0x34, 0xD3, 0x99)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)

FONT = "Segoe UI"
FONT_SEMI = "Segoe UI Semibold"
FONT_LIGHT = "Segoe UI Light"

TOTAL = 20


def _hex(color: RGBColor) -> str:
    return "%02X%02X%02X" % (color[0], color[1], color[2])


def set_alpha(shape, opacity: float) -> None:
    """Set solid-fill opacity (0..1)."""
    solid = shape.fill._xPr.find(qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(qn("a:srgbClr"))
    if clr is None:
        return
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    alpha = etree.SubElement(clr, qn("a:alpha"))
    alpha.set("val", str(int(max(0.0, min(1.0, opacity)) * 100000)))


def no_line(shape) -> None:
    shape.line.fill.background()


def panel(slide, x, y, w, h, fill=PANEL, line=GRID, radius=0.05, accent=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        no_line(sh)
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    if accent is not None:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.02), Inches(y), Inches(w - 0.04), Inches(0.045)
        )
        bar.adjustments[0] = 0.5
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        no_line(bar)
        bar.shadow.inherit = False
    return sh


def tx(
    slide,
    x,
    y,
    w,
    h,
    paras,
    size=11,
    color=TEXT,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line=1.1,
    space=0,
    italic=False,
    margin=0.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor

    if isinstance(paras, str):
        paras = [paras]

    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_before = Pt(0)
        p.space_after = Pt(space)
        runs = para if isinstance(para, list) else [(para, {})]
        for text, ov in runs:
            r = p.add_run()
            r.text = text
            r.font.name = ov.get("font", font)
            r.font.size = Pt(ov.get("size", size))
            r.font.bold = ov.get("bold", bold)
            r.font.italic = ov.get("italic", italic)
            r.font.color.rgb = ov.get("color", color)
            if "baseline" in ov:
                r.font._rPr.set("baseline", str(ov["baseline"]))
    return box


def header(slide, num, section, title, sub=None, accent=CYAN):
    tx(slide, M, 0.34, 8.0, 0.24, section.upper(), size=8.5, color=accent, bold=True)
    tx(slide, M, 0.55, SW - 2 * M, 0.62, title, size=26, color=TEXT, bold=True, font=FONT_SEMI, line=1.0)
    if sub:
        tx(slide, M, 1.14, SW - 2 * M, 0.34, sub, size=10.5, color=MUTED, line=1.05)
    tx(slide, SW - M - 1.4, 0.34, 1.4, 0.24, f"{num:02d} / {TOTAL}", size=8.5, color=DIM, align=PP_ALIGN.RIGHT)


def footer(slide, num, accent=CYAN):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(M), Inches(6.98), Inches(SW - M), Inches(6.98)
    )
    line.line.color.rgb = GRID
    line.line.width = Pt(0.75)
    tx(slide, M, 7.08, 6.0, 0.22, "BURJ KHALIFA  ·  WIND ENGINEERING DOSSIER", size=7.5, color=DIM)
    tx(slide, SW - M - 2.0, 7.08, 2.0, 0.22, f"{num:02d}", size=8, color=accent, bold=True, align=PP_ALIGN.RIGHT)


def source(slide, text, y=6.62):
    tx(slide, M, y, SW - 2 * M, 0.3, text, size=7.2, color=DIM, line=1.05)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def chip(slide, x, y, w, text, color=CYAN, fill=None, size=8.5):
    panel(slide, x, y, w, 0.3, fill=fill or PANEL2, line=None, radius=0.3)
    tx(slide, x, y + 0.02, w, 0.26, text, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def kpi(slide, x, y, w, h, value, unit, label, accent=CYAN):
    panel(slide, x, y, w, h, fill=PANEL, line=GRID)
    tx(slide, x + 0.16, y + 0.14, w - 0.3, 0.6, [[(value, {"size": 27, "bold": True, "color": TEXT}), (unit, {"size": 11, "color": accent, "bold": True})]], line=0.95)
    tx(slide, x + 0.16, y + h - 0.5, w - 0.3, 0.4, label, size=7.6, color=MUTED, line=1.02)


def arrow(slide, x, y, w, h, color, rot=0, opacity=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.rotation = rot
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if opacity < 1:
        set_alpha(sh, opacity)
    no_line(sh)
    sh.shadow.inherit = False
    return sh


def line(slide, x1, y1, x2, y2, color=GRID, width=1.0, dash=None, opacity=1.0):
    sh = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    if dash is not None:
        sh.line.dash_style = dash
    if opacity < 1:
        solid = sh.line.color._xFill
        c = solid.find(qn("a:srgbClr"))
        if c is not None:
            a = etree.SubElement(c, qn("a:alpha"))
            a.set("val", str(int(opacity * 100000)))
    return sh


def polyline(slide, pts, color=CYAN, width=1.0, opacity=1.0, close=False, fill=None):
    q = [(round(x * 100), round(y * 100)) for x, y in pts]
    fb = slide.shapes.build_freeform(q[0][0], q[0][1], scale=9144)
    fb.add_line_segments(q[1:], close=close)
    sh = fb.convert_to_shape()
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if opacity < 1:
            set_alpha(sh, opacity)
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    sh.shadow.inherit = False
    return sh


def wing_plan(slide, cx, cy, radius, width, angle, fill=None, line_color=CYAN, lw=1.1, opacity=0.12):
    th = math.radians(angle)
    ccx = cx + radius / 2 * math.cos(th)
    ccy = cy + radius / 2 * math.sin(th)
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(ccx - radius / 2),
        Inches(ccy - width / 2),
        Inches(radius),
        Inches(width),
    )
    sh.rotation = angle
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        set_alpha(sh, opacity)
    sh.line.color.rgb = line_color
    sh.line.width = Pt(lw)
    sh.shadow.inherit = False


def y_plan(slide, cx, cy, radius, width, rot=0, fill=PANEL2, line_color=CYAN, opacity=0.25, core=True, lw=1.1):
    for a in (rot, rot + 120, rot + 240):
        wing_plan(slide, cx, cy, radius, width, a, fill=fill, line_color=line_color, lw=lw, opacity=opacity)
    if core:
        r = width * 0.44
        sh = slide.shapes.add_shape(
            MSO_SHAPE.HEXAGON, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r)
        )
        sh.rotation = 30
        sh.fill.solid()
        sh.fill.fore_color.rgb = PANEL
        sh.line.color.rgb = line_color
        sh.line.width = Pt(lw)
        sh.shadow.inherit = False


# ---------------------------------------------------------------- charts
def _transparent(chart):
    cs = chart._chartSpace
    chart_el = cs.find(qn("c:chart"))
    if chart_el is None:
        return
    sp = etree.SubElement(cs, qn("c:spPr"))
    chart_el.addnext(sp)
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def prep_chart(chart, legend=False):
    chart.has_title = False
    chart.has_legend = legend
    chart.font.name = FONT
    chart.font.size = Pt(9)
    chart.font.color.rgb = MUTED
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(8.5)
        chart.legend.font.name = FONT
        chart.legend.font.color.rgb = MUTED
    _transparent(chart)


def axis_style(ax, grid=False, fmt="0"):
    ax.has_major_gridlines = grid
    if grid:
        ax.major_gridlines.format.line.color.rgb = GRID
        ax.major_gridlines.format.line.width = Pt(0.5)
    ax.format.line.color.rgb = GRID
    ax.format.line.width = Pt(0.75)
    ax.major_tick_mark = XL_TICK_MARK.NONE
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    ax.tick_labels.font.size = Pt(8.5)
    ax.tick_labels.font.name = FONT
    ax.tick_labels.font.color.rgb = MUTED
    ax.tick_labels.number_format = fmt
    ax.tick_labels.number_format_is_linked = False


def low_tick_labels(ax):
    """Push the axis tick labels to the low edge (useful when the axis crosses at zero)."""
    pos = ax._element.find(qn("c:tickLblPos"))
    if pos is not None:
        pos.set("val", "low")


def bar_chart(slide, x, y, w, h, cats, vals, colors, vmax=None, vmin=0, unit=None, gap=55, fmt="0", labels=True):
    data = CategoryChartData()
    data.categories = cats
    data.add_series("value", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    prep_chart(ch)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    ser = ch.series[0]
    for i, color in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = color
        pt.format.line.fill.background()
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(13)
        dl.font.bold = True
        dl.font.name = FONT_SEMI
        dl.font.color.rgb = TEXT
        dl.number_format = fmt
        dl.number_format_is_linked = False
    ax = ch.value_axis
    axis_style(ax, grid=True, fmt=fmt)
    ax.minimum_scale = vmin
    if vmax is not None:
        ax.maximum_scale = vmax
    if unit is not None:
        ax.major_unit = unit
    axis_style(ch.category_axis, grid=False)
    ch.category_axis.tick_labels.font.size = Pt(8.5)
    return ch


def scatter_chart(slide, x, y, w, h, series, xmax, ymax, xunit=None, yunit=None, xmin=0, ymin=0, xfmt="0", yfmt="0"):
    data = XyChartData()
    for spec in series:
        s = data.add_series(spec["name"])
        for px, py in spec["points"]:
            s.add_data_point(px, py)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    prep_chart(ch, legend=len(series) > 1)
    for spec, s in zip(series, ch.series):
        s.format.line.color.rgb = spec["color"]
        s.format.line.width = Pt(spec.get("width", 2.4))
        if spec.get("dash"):
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        s.smooth = spec.get("smooth", False)
        if spec.get("marker", True):
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = spec.get("marker_size", 6)
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = spec["color"]
            s.marker.format.line.color.rgb = BG
            s.marker.format.line.width = Pt(1.2)
        else:
            s.marker.style = XL_MARKER_STYLE.NONE
    xa = ch.category_axis
    axis_style(xa, grid=False, fmt=xfmt)
    xa.minimum_scale = xmin
    xa.maximum_scale = xmax
    if xunit:
        xa.major_unit = xunit
    ya = ch.value_axis
    axis_style(ya, grid=True, fmt=yfmt)
    ya.minimum_scale = ymin
    ya.maximum_scale = ymax
    if yunit:
        ya.major_unit = yunit
    return ch


# ---------------------------------------------------------------- tables
def table(slide, x, y, w, h, rows, widths, header_size=8.8, body_size=8.6, row_h=0.52, header_h=0.46, accent=CYAN):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, len(rows)):
        tbl.rows[r].height = Inches(row_h)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL2 if r == 0 else (PANEL if r % 2 else BG2)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.08)
            tf.margin_right = Inches(0.06)
            tf.margin_top = Inches(0.02)
            tf.margin_bottom = Inches(0.02)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.0
            run = p.add_run()
            run.text = text
            run.font.name = FONT_SEMI if r == 0 else FONT
            run.font.size = Pt(header_size if r == 0 else body_size)
            run.font.bold = r == 0
            run.font.color.rgb = accent if r == 0 else TEXT
    return tbl
